"""Tests for DOCX document extraction."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image

from services.document_import.extract import extract_document
from services.document_import.models import image_cap_warning


def _make_docx(path: Path) -> None:
    """Build a tiny DOCX with text, table, and embedded image."""
    document = Document()
    document.add_paragraph("Intro paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Method"
    table.cell(0, 1).text = "Path"
    table.cell(1, 0).text = "GET"
    table.cell(1, 1).text = "/users"

    document.add_picture(_png_stream("red"))
    document.save(str(path))


def test_docx_text_and_table_blocks(tmp_path: Path) -> None:
    """DOCX paragraphs and tables become ordered blocks."""
    docx_path = tmp_path / "sample.docx"
    _make_docx(docx_path)
    doc = extract_document(str(docx_path), run_ocr=False)
    kinds = [block["kind"] for block in doc["blocks"]]
    assert "text" in kinds
    assert "table" in kinds
    text_blocks = [block for block in doc["blocks"] if block["kind"] == "text"]
    assert any("Intro paragraph" in block["text"] for block in text_blocks)
    table_blocks = [block for block in doc["blocks"] if block["kind"] == "table"]
    assert table_blocks
    assert "GET" in table_blocks[0]["text"]


def test_docx_embedded_image_block(tmp_path: Path) -> None:
    """DOCX embedded media becomes image blocks."""
    docx_path = tmp_path / "sample.docx"
    _make_docx(docx_path)
    doc = extract_document(str(docx_path), run_ocr=False)
    assert doc["images"]
    assert any(block["kind"] == "image" for block in doc["blocks"])


def test_docx_blocks_preserve_document_order(tmp_path: Path) -> None:
    """Paragraphs, tables, and inline images stay in body order."""
    docx_path = tmp_path / "ordered.docx"
    _make_docx(docx_path)
    doc = extract_document(str(docx_path), run_ocr=False)
    assert [block["kind"] for block in doc["blocks"]] == ["text", "table", "image"]


def _png_stream(color: str, *, size: int = 96) -> BytesIO:
    """Return an in-memory PNG large enough to survive the page-furniture filter."""
    buf = BytesIO()
    Image.new("RGB", (size, size), color=color).save(buf, format="PNG")
    buf.seek(0)
    return buf


def test_docx_extracts_table_cell_image(tmp_path: Path) -> None:
    """Screenshots inside table cells are extracted next to their table."""
    docx_path = tmp_path / "cells.docx"
    document = Document()
    document.add_paragraph("Endpoint: GET /users")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Example response"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(_png_stream("red"))
    document.add_paragraph("Endpoint: POST /users")
    document.paragraphs[-1].add_run().add_picture(_png_stream("blue"))
    document.save(str(docx_path))

    doc = extract_document(str(docx_path), run_ocr=False)
    assert len(doc["images"]) == 2
    kinds = [block["kind"] for block in doc["blocks"]]
    assert kinds == ["text", "table", "image", "text", "image"]
    assert doc["warnings"] == []


def test_docx_extracts_header_image_as_unpositioned(tmp_path: Path) -> None:
    """Header/footer images are still returned, flagged as unpositioned."""
    docx_path = tmp_path / "header.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].add_run().add_picture(_png_stream("green"))
    document.add_paragraph("Body text")
    document.save(str(docx_path))

    doc = extract_document(str(docx_path), run_ocr=False)
    assert len(doc["images"]) == 1
    assert doc["blocks"][-1]["kind"] == "image"
    assert any("not positioned in the document flow" in w for w in doc["warnings"])


def test_docx_image_cap_warns_once(tmp_path: Path) -> None:
    """Hitting max_images emits a single warning across paragraphs."""
    docx_path = tmp_path / "many.docx"
    document = Document()
    for index in range(4):
        document.add_paragraph(f"Shot {index}")
        document.paragraphs[-1].add_run().add_picture(_png_stream("red" if index % 2 else "blue"))
    document.save(str(docx_path))

    doc = extract_document(str(docx_path), max_images=1, run_ocr=False)
    assert len(doc["images"]) == 1
    assert doc["warnings"].count(image_cap_warning(1)) == 1


def test_docx_private_api_contract() -> None:
    """Fail loudly if python-docx internals used for image walking move."""
    from docx.oxml.ns import qn
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    document = Document()
    paragraph = document.add_paragraph("x")
    assert hasattr(paragraph, "_element")
    assert hasattr(document.element, "body")
    assert hasattr(document.part, "related_parts")
    assert qn("a:blip").startswith("{")
    assert CT_P is not None
    assert CT_Tbl is not None
